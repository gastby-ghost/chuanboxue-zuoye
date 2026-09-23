#!/usr/bin/env python3
"""把 Markdown 转成 Word（.docx）。

用法：
    pip install python-docx
    python3 tools/md2docx.py 研究总览.md 研究总览.docx

说明：本仓库以 .md 为版本化的源文件，.docx 是生成物（已在 .gitignore 中忽略）。
需要重做 Word 时跑一次即可，不必把二进制塞进 git 历史。

支持的语法：标题（#–####）、表格（| |）、无序/有序列表、引用（>）、
代码块（```）、粗体（**）、行内代码（`）、图片（![alt](path)）。
另有本仓库的简化写法：「（配图见 `xxx.png`）」会直接插入该图片。
"""
import re
import sys
import os

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("缺少依赖：请先运行  pip install python-docx")

CN_FONT = "STHeiti"          # 中文字体（macOS 自带；Windows 可改为「微软雅黑」）
MONO_FONT = "Courier New"
H_SIZE = {1: 17, 2: 14, 3: 12, 4: 11}


def add_runs(p, text):
    """把一行文本写进段落，处理 **粗体** 与 `行内代码`。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(9.5)
        else:
            r = p.add_run(part)
        r.font.name = r.font.name or CN_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def is_sep_row(cells):
    return all(set(c) <= set("-: ") for c in cells)


def convert(src, out):
    md = open(src, encoding="utf-8").read()
    md = md.replace("&#40;", "(").replace("&#41;", ")")
    base = os.path.dirname(os.path.abspath(src))

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = CN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)

    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]

        if ln.strip() in ("---", ""):
            i += 1
            continue

        # 代码块
        if ln.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run("\n".join(buf))
            r.font.name = MONO_FONT
            r.font.size = Pt(8.5)
            continue

        # 图片：本仓库约定「（配图见 `x.png`）」或标准 ![alt](path)
        m_img = re.search(r"!\[[^\]]*\]\(([^)]+)\)", ln)
        if m_img:
            path = m_img.group(1)
            if not os.path.isabs(path):
                path = os.path.join(base, path)
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(5.0))
            i += 1
            continue
        m_pic = re.search(r"配图见\s*`([^`]+\.(?:png|jpg|jpeg))`", ln)
        if m_pic:
            path = os.path.join(base, m_pic.group(1))
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(5.0))
            i += 1
            continue

        # 表格
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in rows if not is_sep_row(r)]
            if rows:
                ncol = len(rows[0])
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row[:ncol]):
                        cp = t.cell(ri, ci).paragraphs[0]
                        add_runs(cp, cell)
                        for r in cp.runs:
                            r.font.size = Pt(9)
                            if ri == 0:
                                r.bold = True
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            lvl = min(len(m.group(1)), 4)
            h = doc.add_heading(level=lvl)
            r = h.add_run(m.group(2))
            r.font.name = CN_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
            r.font.size = Pt(H_SIZE[lvl])
            r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # 引用
        if ln.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, ln.lstrip("> ").strip())
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 列表
        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", ln))
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", ln):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\.\s+", "", ln))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, ln)
        i += 1

    doc.save(out)
    print("已生成 %s（%d 段落，%d 表格）" % (out, len(doc.paragraphs), len(doc.tables)))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    convert(sys.argv[1], sys.argv[2])
